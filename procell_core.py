"""Core analysis engine for Procell Excel mapping quality checks."""
from __future__ import annotations

from collections import defaultdict
from io import BytesIO
import re
from typing import Any, Iterable

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

APP_NAME = "Procell"
MANDATORY_TARGET_COLUMNS = [
    "USUBJID", "SUBJID", "SPID", "INVID", "VISIT", "VISITNUM",
    "DCMDATE", "GRPID", "STUDYID", "DOMAIN", "CVISIT",
]
DUPLICATE_KEY_COLUMNS = ["Target Table", "Target Column", "Pseudo Logic", "Transpose Rule"]
REQUIRED_COLUMNS = ["Source Table", "Target Table", "Target Column", "Pseudo Logic", "Transpose Rule"]


def _header_key(value: Any) -> str:
    """Normalize header punctuation/case so common variants still match."""
    return re.sub(r"[^a-z0-9]", "", str(value).strip().casefold())


_CANONICAL = {
    "sourcetable": "Source Table",
    "targettable": "Target Table",
    "targetcolumn": "Target Column",
    "pseudologic": "Pseudo Logic",
    "transposerule": "Transpose Rule",
}


def _is_blank(value: Any) -> bool:
    if value is None:
        return True
    try:
        if pd.isna(value):
            return True
    except (TypeError, ValueError):
        pass
    return isinstance(value, str) and not value.strip()


def _display(value: Any) -> str:
    return "" if _is_blank(value) else str(value).strip()


def _key_value(value: Any) -> str:
    """Comparison value: blanks compare as blanks; text comparison is case-insensitive."""
    if _is_blank(value):
        return ""
    if isinstance(value, str):
        return re.sub(r"\s+", " ", value.strip()).casefold()
    return str(value).strip().casefold()


def load_sheet_from_bytes(content: bytes, sheet_name: str) -> tuple[pd.DataFrame, list[int]]:
    """Read an XLSX/XLSM worksheet while preserving physical Excel row numbers.

    The first row is treated as the header row. Completely blank data rows are
    ignored, but their physical positions are retained for subsequent rows.
    """
    try:
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=False)
    except Exception as exc:
        raise ValueError(f"Could not open this workbook. Please upload a valid .xlsx or .xlsm file. Details: {exc}") from exc

    try:
        if sheet_name not in workbook.sheetnames:
            raise ValueError(f"Worksheet '{sheet_name}' was not found in the uploaded workbook.")
        ws = workbook[sheet_name]
        row_iter = ws.iter_rows(values_only=True)
        header_row = next(row_iter, None)
        if header_row is None:
            raise ValueError("The selected worksheet is empty.")

        headers: list[str] = []
        seen: dict[str, int] = defaultdict(int)
        for idx, value in enumerate(header_row, start=1):
            label = str(value).strip() if value is not None and str(value).strip() else f"Unnamed: {idx}"
            seen[label.casefold()] += 1
            if seen[label.casefold()] > 1:
                label = f"{label} ({seen[label.casefold()]})"
            headers.append(label)

        data: list[list[Any]] = []
        excel_rows: list[int] = []
        for excel_row_number, row in enumerate(row_iter, start=2):
            values = list(row[:len(headers)])
            if len(values) < len(headers):
                values.extend([None] * (len(headers) - len(values)))
            if all(_is_blank(v) for v in values):
                continue
            data.append(values)
            excel_rows.append(excel_row_number)
        frame = pd.DataFrame(data, columns=headers)
        return frame, excel_rows
    finally:
        workbook.close()


def _resolve_columns(df: pd.DataFrame) -> dict[str, str]:
    found: dict[str, str] = {}
    for col in df.columns:
        key = _header_key(col)
        canonical = _CANONICAL.get(key)
        if canonical:
            if canonical in found:
                raise ValueError(
                    f"More than one column maps to '{canonical}' (for example, '{found[canonical]}' and '{col}'). "
                    "Please keep only one version of each required header."
                )
            found[canonical] = col
    missing = [name for name in REQUIRED_COLUMNS if name not in found]
    if missing:
        raise ValueError(
            "The selected worksheet is missing required columns: " + ", ".join(missing) + ". "
            "Required headers are Source Table, Target Table, Target Column, Pseudo Logic, and Transpose Rule."
        )
    return found


def analyse_mapping(df: pd.DataFrame, excel_rows: Iterable[int] | None = None) -> dict[str, Any]:
    """Run duplicate detection and mandatory target-column coverage checks."""
    df = df.copy()
    if excel_rows is None:
        source_rows = list(range(2, len(df) + 2))
    else:
        source_rows = list(excel_rows)
    if len(source_rows) != len(df):
        raise ValueError("Internal row tracking mismatch while reading the worksheet.")

    cols = _resolve_columns(df)
    df.insert(0, "Procell Excel Row", source_rows)

    # Group rows by all four duplicate-key fields. Blank values compare as blanks.
    duplicate_groups: dict[tuple[str, ...], list[int]] = defaultdict(list)
    for position, (_, row) in enumerate(df.iterrows()):
        key = tuple(_key_value(row[cols[column]]) for column in DUPLICATE_KEY_COLUMNS)
        duplicate_groups[key].append(position)

    duplicate_positions: set[int] = set()
    duplicate_group_id_by_position: dict[int, str] = {}
    duplicate_group_size_by_position: dict[int, int] = {}
    duplicate_key_display_by_position: dict[int, dict[str, str]] = {}
    duplicate_group_number = 0
    for key, positions in duplicate_groups.items():
        if len(positions) < 2:
            continue
        duplicate_group_number += 1
        group_id = f"DUP-{duplicate_group_number:04d}"
        for position in positions:
            duplicate_positions.add(position)
            duplicate_group_id_by_position[position] = group_id
            duplicate_group_size_by_position[position] = len(positions)
            row = df.iloc[position]
            duplicate_key_display_by_position[position] = {
                column: _display(row[cols[column]]) for column in DUPLICATE_KEY_COLUMNS
            }

    duplicate_positions_sorted = sorted(duplicate_positions, key=lambda p: source_rows[p])
    duplicate_rows = df.iloc[duplicate_positions_sorted].copy() if duplicate_positions_sorted else df.iloc[0:0].copy()
    if len(duplicate_rows):
        duplicate_rows["Procell Duplicate Group"] = [duplicate_group_id_by_position[p] for p in duplicate_positions_sorted]
        duplicate_rows["Procell Group Size"] = [duplicate_group_size_by_position[p] for p in duplicate_positions_sorted]
        for column in DUPLICATE_KEY_COLUMNS:
            duplicate_rows[f"Procell Key - {column}"] = [duplicate_key_display_by_position[p][column] for p in duplicate_positions_sorted]
    else:
        duplicate_rows["Procell Duplicate Group"] = pd.Series(dtype="object")
        duplicate_rows["Procell Group Size"] = pd.Series(dtype="int64")
        for column in DUPLICATE_KEY_COLUMNS:
            duplicate_rows[f"Procell Key - {column}"] = pd.Series(dtype="object")

    # Build source-target table pairs and aggregate the target columns seen per pair.
    pair_records: dict[tuple[str, str], dict[str, Any]] = {}
    skipped_blank_pair_rows = 0
    for position, (_, row) in enumerate(df.iterrows()):
        source_display = _display(row[cols["Source Table"]])
        target_display = _display(row[cols["Target Table"]])
        source_key = _key_value(row[cols["Source Table"]])
        target_key = _key_value(row[cols["Target Table"]])
        if not source_key or not target_key:
            skipped_blank_pair_rows += 1
            continue
        pair_key = (source_key, target_key)
        if pair_key not in pair_records:
            pair_records[pair_key] = {
                "Source Table": source_display,
                "Target Table": target_display,
                "rows": [],
                "target_columns": set(),
            }
        pair_records[pair_key]["rows"].append(position)
        target_col_key = _key_value(row[cols["Target Column"]])
        if target_col_key:
            pair_records[pair_key]["target_columns"].add(target_col_key)

    pair_summary_records: list[dict[str, Any]] = []
    missing_detail_records: list[dict[str, Any]] = []
    missing_columns_by_pair: dict[tuple[str, str], list[str]] = {}
    pair_source_keys: dict[tuple[str, str], tuple[str, str]] = {}

    for pair_number, (pair_key, info) in enumerate(pair_records.items(), start=1):
        seen = info["target_columns"]
        missing = [col for col in MANDATORY_TARGET_COLUMNS if _key_value(col) not in seen]
        missing_columns_by_pair[pair_key] = missing
        pair_source_keys[pair_key] = (info["Source Table"], info["Target Table"])
        status = "Missing mandatory columns" if missing else "Complete"
        pair_summary_records.append({
            "Source Table": info["Source Table"],
            "Target Table": info["Target Table"],
            "Source-Target Pair Rows": len(info["rows"]),
            "Mandatory Target Columns Present": len(MANDATORY_TARGET_COLUMNS) - len(missing),
            "Mandatory Target Columns Required": len(MANDATORY_TARGET_COLUMNS),
            "Missing Mandatory Count": len(missing),
            "Missing Mandatory Target Columns": ", ".join(missing),
            "Check Status": status,
        })
        for target_column in missing:
            missing_detail_records.append({
                "Source Table": info["Source Table"],
                "Target Table": info["Target Table"],
                "Missing Target Column": target_column,
                "Source-Target Pair Rows": len(info["rows"]),
            })

    pair_summary_columns = [
        "Source Table", "Target Table", "Source-Target Pair Rows",
        "Mandatory Target Columns Present", "Mandatory Target Columns Required",
        "Missing Mandatory Count", "Missing Mandatory Target Columns", "Check Status",
    ]
    missing_detail_columns = ["Source Table", "Target Table", "Missing Target Column", "Source-Target Pair Rows"]
    pair_summary = pd.DataFrame(pair_summary_records, columns=pair_summary_columns)
    missing_mandatory = pd.DataFrame(missing_detail_records, columns=missing_detail_columns)

    # Full input with row-level duplicate flag and pair-level missing-column context.
    checked_input = df.copy()
    checked_input["Procell Duplicate Status"] = ["Duplicate" if p in duplicate_positions else "Unique" for p in range(len(df))]
    checked_input["Procell Duplicate Group"] = [duplicate_group_id_by_position.get(p, "") for p in range(len(df))]
    checked_input["Procell Duplicate Group Size"] = [duplicate_group_size_by_position.get(p, 1) for p in range(len(df))]
    missing_for_row: list[str] = []
    pair_status_for_row: list[str] = []
    for position in range(len(df)):
        row = df.iloc[position]
        pair_key = (_key_value(row[cols["Source Table"]]), _key_value(row[cols["Target Table"]]))
        if pair_key in missing_columns_by_pair:
            missing_for_row.append(", ".join(missing_columns_by_pair[pair_key]))
            pair_status_for_row.append("Missing mandatory columns" if missing_columns_by_pair[pair_key] else "Complete")
        else:
            missing_for_row.append("")
            pair_status_for_row.append("Skipped: blank Source Table or Target Table")
    checked_input["Procell Missing Mandatory Columns for Pair"] = missing_for_row
    checked_input["Procell Pair Check Status"] = pair_status_for_row

    duplicate_group_count = duplicate_group_number
    duplicate_row_count = len(duplicate_positions)
    missing_pair_count = int((pair_summary["Missing Mandatory Count"] > 0).sum()) if len(pair_summary) else 0
    total_missing_mappings = len(missing_mandatory)
    metrics = [
        {"Metric": "Input data rows analyzed", "Value": len(df), "Details": "Non-empty data rows from the selected worksheet."},
        {"Metric": "Duplicate groups", "Value": duplicate_group_count, "Details": "Groups where Target Table, Target Column, Pseudo Logic, and Transpose Rule match."},
        {"Metric": "Rows marked as duplicates", "Value": duplicate_row_count, "Details": "All rows in a repeated group are listed, including the first occurrence."},
        {"Metric": "Source-Target pairs checked", "Value": len(pair_records), "Details": "Pairs with both Source Table and Target Table populated."},
        {"Metric": "Pairs missing mandatory columns", "Value": missing_pair_count, "Details": "Source-Target pairs missing at least one of the 11 mandatory Target Column values."},
        {"Metric": "Total missing mandatory mappings", "Value": total_missing_mappings, "Details": "One row per missing mandatory Target Column per Source-Target pair."},
        {"Metric": "Rows skipped from pair check", "Value": skipped_blank_pair_rows, "Details": "Rows where Source Table or Target Table is blank; still included in duplicate checking."},
    ]
    summary = pd.DataFrame(metrics, columns=["Metric", "Value", "Details"])

    methodology = pd.DataFrame([
        {"Topic": "Duplicate rule", "Description": "Rows are grouped by Target Table + Target Column + Pseudo Logic + Transpose Rule. Every row in a group of size 2 or more is marked as a duplicate, including the first occurrence."},
        {"Topic": "Comparison behavior", "Description": "Header names tolerate spacing/punctuation/case variations. Text values are trimmed and compared case-insensitively; repeated internal whitespace is collapsed for matching. Original cell values remain unchanged in the report."},
        {"Topic": "Missing mandatory rule", "Description": "Each populated Source Table + Target Table pair is checked for all 11 mandatory target columns. Matching Target Column values are case-insensitive and trimmed."},
        {"Topic": "Mandatory target columns", "Description": ", ".join(MANDATORY_TARGET_COLUMNS)},
        {"Topic": "Blank table names", "Description": "Rows with a blank Source Table or Target Table are excluded from pair coverage checks, but remain in the duplicate check and Checked Input sheet."},
        {"Topic": "Row numbering", "Description": "Procell Excel Row records the physical worksheet row. The header is assumed to be in row 1."},
        {"Topic": "Empty result sheets", "Description": "If no duplicates or mandatory columns are missing, the corresponding detail sheet contains headers only."},
    ])

    return {
        "summary": summary,
        "duplicate_rows": duplicate_rows,
        "missing_mandatory": missing_mandatory,
        "pair_summary": pair_summary,
        "checked_input": checked_input,
        "methodology": methodology,
        "metrics": {row["Metric"]: row["Value"] for row in metrics},
    }


def build_report_xlsx(results: dict[str, Any]) -> bytes:
    """Create a styled multi-sheet XLSX report and return its bytes."""
    output = BytesIO()
    sheet_frames = [
        ("Summary", results["summary"]),
        ("Duplicate Rows", results["duplicate_rows"]),
        ("Missing Mandatory", results["missing_mandatory"]),
        ("Pair Summary", results["pair_summary"]),
        ("Checked Input", results["checked_input"]),
        ("Methodology", results["methodology"]),
    ]
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        for sheet_name, frame in sheet_frames:
            frame.to_excel(writer, sheet_name=sheet_name, index=False)

    output.seek(0)
    workbook = load_workbook(output)
    navy = "183153"
    blue = "DCEAF7"
    pale_red = "FCE8E6"
    pale_green = "E2F0D9"
    white = "FFFFFF"
    gray = "D9E2F3"
    thin_gray = Side(style="thin", color=gray)

    for ws in workbook.worksheets:
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions if ws.max_row >= 1 and ws.max_column >= 1 else None
        ws.sheet_view.showGridLines = False
        for cell in ws[1]:
            cell.fill = PatternFill("solid", fgColor=navy)
            cell.font = Font(name="Aptos", size=10, bold=True, color=white)
            cell.alignment = Alignment(vertical="center", wrap_text=True)
            cell.border = Border(bottom=Side(style="medium", color="6E9BC5"))
        ws.row_dimensions[1].height = 32
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.font = Font(name="Aptos", size=10, color="1F2937")
                cell.alignment = Alignment(vertical="top", wrap_text=True)
                if cell.row % 2 == 0:
                    cell.fill = PatternFill("solid", fgColor="F7FAFC")
                cell.border = Border(bottom=thin_gray)
        for col_idx in range(1, ws.max_column + 1):
            values = [str(ws.cell(row=r, column=col_idx).value or "") for r in range(1, min(ws.max_row, 250) + 1)]
            max_len = max([len(value) for value in values] + [10])
            width = min(max(max_len + 2, 12), 42)
            ws.column_dimensions[get_column_letter(col_idx)].width = width
        if ws.max_row > 1:
            ws.row_dimensions[1].height = 36

    # Emphasize status columns where present.
    for name in ("Pair Summary", "Checked Input"):
        ws = workbook[name]
        headers = {cell.value: cell.column for cell in ws[1]}
        status_col = headers.get("Check Status") or headers.get("Procell Pair Check Status")
        if status_col:
            for row in range(2, ws.max_row + 1):
                cell = ws.cell(row=row, column=status_col)
                value = str(cell.value or "")
                if "Missing" in value or "Skipped" in value:
                    cell.fill = PatternFill("solid", fgColor=pale_red)
                elif value == "Complete":
                    cell.fill = PatternFill("solid", fgColor=pale_green)

    # Summary sheet presentation.
    summary_ws = workbook["Summary"]
    summary_ws.freeze_panes = "A2"
    summary_ws.column_dimensions["A"].width = 34
    summary_ws.column_dimensions["B"].width = 16
    summary_ws.column_dimensions["C"].width = 82
    summary_ws.auto_filter.ref = summary_ws.dimensions

    workbook.properties.title = "Procell Mapping Quality Report"
    workbook.properties.subject = "Duplicate mapping and mandatory target column coverage checks"
    workbook.properties.creator = "Procell"

    final_output = BytesIO()
    workbook.save(final_output)
    workbook.close()
    return final_output.getvalue()
