"""Streamlit user interface for Procell."""
from __future__ import annotations

from io import BytesIO
from pathlib import Path

import pandas as pd
import streamlit as st
from openpyxl import load_workbook

from procell_core import analyse_mapping, build_report_xlsx, load_sheet_from_bytes, MANDATORY_TARGET_COLUMNS

st.set_page_config(page_title="Procell | Excel Mapping Checker", page_icon="🧩", layout="wide")

st.markdown(
    """
    <style>
    .block-container {padding-top: 2.1rem; padding-bottom: 2.5rem; max-width: 1250px;}
    .procell-header {padding: 1.15rem 1.4rem; border-radius: 14px; background: linear-gradient(115deg, #183153 0%, #315f86 100%); color: white; margin-bottom: 1.15rem;}
    .procell-header h1 {color: white; margin: 0; font-size: 2.2rem;}
    .procell-header p {color: #e7f0f7; margin: .35rem 0 0 0; font-size: 1rem;}
    .section-note {color: #52616f; font-size: .93rem;}
    </style>
    """,
    unsafe_allow_html=True,
)

st.markdown(
    '<div class="procell-header"><h1>Procell</h1><p>Excel mapping quality checker · Duplicate detection and mandatory target-column coverage</p></div>',
    unsafe_allow_html=True,
)

with st.expander("What Procell checks", expanded=False):
    st.markdown(
        "**1. Duplicate mapping rows** are grouped by `Target Table`, `Target Column`, `Pseudo Logic`, and `Transpose Rule`. "
        "Every row in a repeated group is reported, including the first occurrence.\n\n"
        "**2. Mandatory target columns** are checked for every populated `Source Table` + `Target Table` pair. "
        "The required columns are: " + ", ".join(MANDATORY_TARGET_COLUMNS) + ".\n\n"
        "Header names and values are matched without case sensitivity; leading/trailing spaces are ignored. "
        "The header row must be the first row in the selected worksheet."
    )

uploaded_file = st.file_uploader("Upload an Excel mapping workbook", type=["xlsx", "xlsm"], help="Supported formats: .xlsx and .xlsm")

if uploaded_file is None:
    st.info("Upload a workbook to begin. Your file is processed in this app session; no external service is used by the Procell code.")
    st.markdown("**Expected headers include:** `Source Table`, `Target Table`, `Target Column`, `Pseudo Logic`, and `Transpose Rule`. Other columns are preserved in the output.")
    st.stop()

file_bytes = uploaded_file.getvalue()
try:
    workbook = load_workbook(BytesIO(file_bytes), read_only=True, data_only=False)
    sheet_names = workbook.sheetnames
    workbook.close()
except Exception as exc:
    st.error(f"Could not open this Excel workbook. Please check that the file is valid and is not password-protected. Details: {exc}")
    st.stop()

selected_sheet = st.selectbox("Choose the worksheet to analyze", options=sheet_names, index=0)

try:
    preview_df, _preview_rows = load_sheet_from_bytes(file_bytes, selected_sheet)
    st.caption(f"Selected worksheet: **{selected_sheet}** · {len(preview_df):,} non-empty data row(s) detected · header row assumed to be row 1")
    with st.expander("Preview input data", expanded=False):
        st.dataframe(preview_df.head(15), use_container_width=True, hide_index=True)
except Exception as exc:
    st.error(str(exc))
    st.stop()

if st.button("Run Procell checks", type="primary", use_container_width=False):
    try:
        input_df, source_rows = load_sheet_from_bytes(file_bytes, selected_sheet)
        results = analyse_mapping(input_df, source_rows)
        report_bytes = build_report_xlsx(results)
        st.session_state["procell_results"] = results
        st.session_state["procell_report_bytes"] = report_bytes
        st.session_state["procell_source_name"] = uploaded_file.name
        st.session_state["procell_sheet_name"] = selected_sheet
        st.success("Checks completed. Review the findings below and download the Excel report.")
    except Exception as exc:
        st.session_state.pop("procell_results", None)
        st.session_state.pop("procell_report_bytes", None)
        st.error(f"Procell could not analyze this worksheet: {exc}")

if "procell_results" in st.session_state:
    results = st.session_state["procell_results"]
    metrics = results["metrics"]
    st.subheader("Results overview")
    metric_cols = st.columns(4)
    metric_cols[0].metric("Rows analyzed", f"{metrics['Input data rows analyzed']:,}")
    metric_cols[1].metric("Duplicate groups", f"{metrics['Duplicate groups']:,}")
    metric_cols[2].metric("Duplicate rows", f"{metrics['Rows marked as duplicates']:,}")
    metric_cols[3].metric("Pairs with missing columns", f"{metrics['Pairs missing mandatory columns']:,}")
    st.caption(f"Total missing mandatory mappings: **{metrics['Total missing mandatory mappings']:,}** · Source-target pairs checked: **{metrics['Source-Target pairs checked']:,}**")

    tabs = st.tabs(["Summary", "Duplicate Rows", "Missing Mandatory", "Pair Summary"])
    with tabs[0]:
        st.dataframe(results["summary"], use_container_width=True, hide_index=True)
    with tabs[1]:
        if results["duplicate_rows"].empty:
            st.success("No duplicate groups found.")
        else:
            st.warning(f"Found {metrics['Rows marked as duplicates']:,} row(s) across {metrics['Duplicate groups']:,} duplicate group(s).")
            st.dataframe(results["duplicate_rows"], use_container_width=True, hide_index=True)
    with tabs[2]:
        if results["missing_mandatory"].empty:
            st.success("All checked Source Table–Target Table pairs contain all 11 mandatory target columns.")
        else:
            st.warning(f"Found {metrics['Total missing mandatory mappings']:,} missing target-column mapping(s).")
            st.dataframe(results["missing_mandatory"], use_container_width=True, hide_index=True)
    with tabs[3]:
        st.dataframe(results["pair_summary"], use_container_width=True, hide_index=True)

    source_name = st.session_state.get("procell_source_name", "mapping.xlsx")
    source_stem = Path(source_name).stem or "mapping"
    output_name = f"{source_stem}_Procell_Report.xlsx"
    st.download_button(
        label="Download consolidated Excel report",
        data=st.session_state["procell_report_bytes"],
        file_name=output_name,
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        type="primary",
    )
    st.caption("The report contains Summary, Duplicate Rows, Missing Mandatory, Pair Summary, Checked Input, and Methodology worksheets.")
